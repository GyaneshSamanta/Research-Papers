import os
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

def set_number_of_columns(section, num_cols):
    sectPr = section._sectPr
    cols = sectPr.xpath('./w:cols')[0]
    cols.set(qn('w:num'), str(num_cols))
    cols.set(qn('w:space'), '720')  # 0.5 inch

def insert_paragraph_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    return new_para

def insert_table_after_paragraph(doc, p, df):
    table = doc.add_table(rows=df.shape[0]+1, cols=df.shape[1])
    table.style = 'Table Grid'
    
    # Headers
    hdr_cells = table.rows[0].cells
    for j, col_name in enumerate(df.columns):
        hdr_cells[j].text = str(col_name)
        hdr_cells[j].paragraphs[0].runs[0].bold = True
        
    # Data
    for i in range(df.shape[0]):
        row_cells = table.rows[i+1].cells
        for j in range(df.shape[1]):
            row_cells[j].text = str(df.iloc[i, j])
            
    # Move the table to be right after the specified paragraph
    p._p.addnext(table._tbl)
    return table

def format_document(doc_path, out_path, visuals_dir):
    doc = Document(doc_path)
    
    # 1. Reformat to 1-column layout
    for section in doc.sections:
        set_number_of_columns(section, 1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Visuals mapping
    visual_map = {
        'RQ1_Fig1_KDE_Gap.png': "Figure 1: KDE Distribution Gap between Consumer and B2B Satisfaction",
        'RQ1_Fig2_ECDF_Lag.png': "Figure 2: Cumulative Distribution of Performance Lag",
        'RQ1_Fig3_ROC_Curve.png': "Figure 3: ROC Curve Predicting Platform Type from Score",
        'RQ2_Fig1_Violin_Efficiency.png': "Figure 4: Violin Plot showing Task Efficiency by Design Type",
        'RQ2_Fig2_Boxplot_Adoption.png': "Figure 5: Engagement Score vs. Feature Adoption Status",
        'RQ3_Fig1_Survival_Curve.png': "Figure 6: Kaplan-Meier Survival Analysis of Adopters vs Non-Adopters",
        'RQ3_Fig2_Scatter_Sales.png': "Figure 7: Regional Feature Adoption vs. Total Sales"
    }
    
    inserted = {img: False for img in visual_map.keys()}
    tables_inserted = {'rq1': False, 'rq2': False, 'rq3': False}
    
    for p in list(doc.paragraphs):
        text_lower = p.text.lower().strip()
        
        # -------------------- EDA SECTIONS (FIGURES) --------------------
        if text_lower.startswith("8.1 eda for h1"):
            cursor_p = p
            for img in ['RQ1_Fig1_KDE_Gap.png', 'RQ1_Fig2_ECDF_Lag.png', 'RQ1_Fig3_ROC_Curve.png']:
                if not inserted[img] and os.path.exists(os.path.join(visuals_dir, img)):
                    # Insert image
                    img_p = insert_paragraph_after(cursor_p)
                    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_p.add_run().add_picture(os.path.join(visuals_dir, img), width=Inches(6.0))
                    cursor_p = img_p
                    
                    # Insert caption
                    cap_p = insert_paragraph_after(cursor_p)
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_run = cap_p.add_run(visual_map[img])
                    cap_run.italic = True
                    cursor_p = cap_p
                    inserted[img] = True
                    
        elif text_lower.startswith("8.2 eda for h2"):
            cursor_p = p
            for img in ['RQ2_Fig1_Violin_Efficiency.png', 'RQ2_Fig2_Boxplot_Adoption.png']:
                if not inserted[img] and os.path.exists(os.path.join(visuals_dir, img)):
                    img_p = insert_paragraph_after(cursor_p)
                    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_p.add_run().add_picture(os.path.join(visuals_dir, img), width=Inches(6.0))
                    cursor_p = img_p
                    
                    cap_p = insert_paragraph_after(cursor_p)
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_run = cap_p.add_run(visual_map[img])
                    cap_run.italic = True
                    cursor_p = cap_p
                    inserted[img] = True

        elif text_lower.startswith("8.3 eda for h3a"):
            cursor_p = p
            for img in ['RQ3_Fig1_Survival_Curve.png']:
                if not inserted[img] and os.path.exists(os.path.join(visuals_dir, img)):
                    img_p = insert_paragraph_after(cursor_p)
                    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_p.add_run().add_picture(os.path.join(visuals_dir, img), width=Inches(6.0))
                    cursor_p = img_p
                    
                    cap_p = insert_paragraph_after(cursor_p)
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_run = cap_p.add_run(visual_map[img])
                    cap_run.italic = True
                    cursor_p = cap_p
                    inserted[img] = True
                    
        elif text_lower.startswith("8.4 eda for h3b"):
            cursor_p = p
            for img in ['RQ3_Fig2_Scatter_Sales.png']:
                if not inserted[img] and os.path.exists(os.path.join(visuals_dir, img)):
                    img_p = insert_paragraph_after(cursor_p)
                    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_p.add_run().add_picture(os.path.join(visuals_dir, img), width=Inches(6.0))
                    cursor_p = img_p
                    
                    cap_p = insert_paragraph_after(cursor_p)
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_run = cap_p.add_run(visual_map[img])
                    cap_run.italic = True
                    cursor_p = cap_p
                    inserted[img] = True
                    
        # -------------------- RESULTS SECTIONS (TABLES) --------------------
        elif text_lower.startswith("12.1 results for h1"):
            cursor_p = p
            if not tables_inserted['rq1'] and os.path.exists(f"{visuals_dir}/table_rq1.csv"):
                df = pd.read_csv(f"{visuals_dir}/table_rq1.csv")
                space_p = insert_paragraph_after(cursor_p)
                cursor_p = space_p
                
                title_p = insert_paragraph_after(cursor_p)
                title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title_p.add_run("Table 3: Statistical Summary of the Satisfaction Gap").bold = True
                cursor_p = title_p
                
                table = insert_table_after_paragraph(doc, cursor_p, df)
                blank_p = OxmlElement("w:p")
                table._tbl.addnext(blank_p)
                cursor_p = Paragraph(blank_p, cursor_p._parent)
                tables_inserted['rq1'] = True
                
        elif text_lower.startswith("12.2 results for h2"):
            cursor_p = p
            if not tables_inserted['rq2'] and os.path.exists(f"{visuals_dir}/table_rq2.csv"):
                df = pd.read_csv(f"{visuals_dir}/table_rq2.csv")
                space_p = insert_paragraph_after(cursor_p)
                cursor_p = space_p
                
                title_p = insert_paragraph_after(cursor_p)
                title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title_p.add_run("Table 4: Efficiency Impact of Intuitive Design").bold = True
                cursor_p = title_p
                
                table = insert_table_after_paragraph(doc, cursor_p, df)
                blank_p = OxmlElement("w:p")
                table._tbl.addnext(blank_p)
                cursor_p = Paragraph(blank_p, cursor_p._parent)
                tables_inserted['rq2'] = True
                
        elif text_lower.startswith("12.4 results for h3b"):
            cursor_p = p
            if not tables_inserted['rq3'] and os.path.exists(f"{visuals_dir}/table_rq3.csv"):
                df = pd.read_csv(f"{visuals_dir}/table_rq3.csv")
                space_p = insert_paragraph_after(cursor_p)
                cursor_p = space_p
                
                title_p = insert_paragraph_after(cursor_p)
                title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title_p.add_run("Table 5: Regional Feature Adoption and Sales Summary").bold = True
                cursor_p = title_p
                
                table = insert_table_after_paragraph(doc, cursor_p, df)
                blank_p = OxmlElement("w:p")
                table._tbl.addnext(blank_p)
                cursor_p = Paragraph(blank_p, cursor_p._parent)
                tables_inserted['rq3'] = True

    doc.save(out_path)
    print(f"Document saved to {out_path}")

if __name__ == "__main__":
    doc_path = "working paper/B2B SaaS Research Paper.docx"
    out_path = "working paper/B2B SaaS Research Paper_Final_Formatted.docx"
    visuals_dir = "upgraded_visuals"
    format_document(doc_path, out_path, visuals_dir)
